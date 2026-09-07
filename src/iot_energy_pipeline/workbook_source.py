"""Read registered workbook rows through an exact source-layout contract."""

from __future__ import annotations

from dataclasses import dataclass
from typing import TYPE_CHECKING, Any, Iterator

from .identity import SourceLocator
from .manifests import DatasetManifest
from .source_layout import SourceLayoutRules, WorkbookLayoutRule

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True)
class WorkbookSourceRow:
    """One physical source row before semantic conversion."""

    locator: SourceLocator
    layout_rule: WorkbookLayoutRule
    timestamp_value: object
    meter_id_value: object
    numeric_value: object
    row_values: tuple[Any, ...]


@dataclass(frozen=True)
class WorkbookSourceStart:
    """Validated workbook boundary, emitted even when it has no data rows."""

    relative_path: str
    sheet_name: str
    layout_rule: WorkbookLayoutRule
    path_meter_id: str | None


def _validated_data_sheet(
    workbook: Any,
    relative_path: str,
    layout: WorkbookLayoutRule,
) -> tuple["Worksheet", dict[str, int]]:
    actual_sheets = tuple(sheet.title for sheet in workbook.worksheets)
    declared_sheets = (layout.sheet_name, *layout.ignored_sheet_names)
    if set(actual_sheets) != set(declared_sheets) or len(actual_sheets) != len(
        declared_sheets
    ):
        raise ValueError(
            f"unexpected workbook sheets in {relative_path}: "
            f"expected {list(declared_sheets)!r}, got {list(actual_sheets)!r}"
        )
    worksheet = workbook[layout.sheet_name]
    rows = worksheet.iter_rows(min_row=layout.header_row, values_only=True)
    try:
        header_values = next(rows)
    except StopIteration as exc:
        raise ValueError(
            f"missing declared header row in {relative_path}:{worksheet.title}"
        ) from exc
    headers = tuple(
        "" if value is None else str(value).strip() for value in header_values
    )
    if headers != layout.expected_columns:
        raise ValueError(
            f"source-layout header mismatch in {relative_path}:{worksheet.title}: "
            f"expected {list(layout.expected_columns)!r}, got {list(headers)!r}"
        )
    return worksheet, {header: index for index, header in enumerate(headers)}


def iter_workbook_source_items(
    manifest: DatasetManifest,
    source_layout: SourceLayoutRules,
) -> Iterator[WorkbookSourceStart | WorkbookSourceRow]:
    """Yield each validated workbook boundary followed by its physical rows."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency gate
        raise RuntimeError("dataset extra is required for workbooks") from exc

    if manifest._root is None:
        raise ValueError("manifest must retain a verified source root")
    manifest._validate_files(manifest._root)

    for file_manifest in manifest.files:
        if file_manifest.media_kind != "workbook":
            continue
        resolved = source_layout.resolve(file_manifest.relative_path)
        workbook = load_workbook(
            manifest._root / file_manifest.relative_path,
            read_only=True,
            data_only=True,
        )
        try:
            worksheet, lookup = _validated_data_sheet(
                workbook,
                file_manifest.relative_path,
                resolved.rule,
            )
            yield WorkbookSourceStart(
                relative_path=file_manifest.relative_path,
                sheet_name=worksheet.title,
                layout_rule=resolved.rule,
                path_meter_id=resolved.path_meter_id,
            )
            rows = worksheet.iter_rows(
                min_row=resolved.rule.header_row + 1,
                values_only=True,
            )
            for physical_row, row in enumerate(
                rows, start=resolved.rule.header_row + 1
            ):
                meter_id = (
                    resolved.path_meter_id
                    if resolved.path_meter_id is not None
                    else row[lookup[resolved.rule.meter_id_column]]  # type: ignore[index]
                )
                yield WorkbookSourceRow(
                    locator=SourceLocator(
                        manifest.dataset_version,
                        manifest.source_variant,
                        file_manifest.relative_path,
                        worksheet.title,
                        physical_row,
                    ),
                    layout_rule=resolved.rule,
                    timestamp_value=row[lookup[resolved.rule.timestamp_column]],
                    meter_id_value=meter_id,
                    numeric_value=row[lookup[resolved.rule.value_column]],
                    row_values=tuple(row),
                )
        finally:
            workbook.close()


def iter_workbook_source_rows(
    manifest: DatasetManifest,
    source_layout: SourceLayoutRules,
) -> Iterator[WorkbookSourceRow]:
    """Yield every declared physical data row in manifest path order."""

    for item in iter_workbook_source_items(manifest, source_layout):
        if isinstance(item, WorkbookSourceRow):
            yield item
