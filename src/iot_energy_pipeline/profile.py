"""Dataset profiling for workbook structure and cell statistics."""

from __future__ import annotations

import hashlib
import json
import os
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from re import compile as re_compile
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from openpyxl.cell.cell import Cell
    from openpyxl.worksheet.worksheet import Worksheet

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format

from iot_energy_pipeline.manifests import DatasetManifest, FileManifest

_ISO_TIMESTAMP_PATTERN = re_compile(
    r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}"
    r"(?:\.\d+)?"
    r"(?:Z|[+-]\d{2}:\d{2})?$"
)


@dataclass(frozen=True)
class SheetProfile:
    """Profile statistics for a single worksheet."""

    sheet_title: str
    max_row: int
    max_column: int
    header_texts: tuple[str, ...]
    null_count: int
    datetime_count: int
    date_count: int
    time_count: int
    numeric_cell_count: int
    numeric_text_count: int
    iso_timestamp_text_count: int
    other_text_count: int


@dataclass(frozen=True)
class WorkbookProfile:
    """Profile for a single workbook file."""

    relative_path: str
    sheets: tuple[SheetProfile, ...]


@dataclass(frozen=True)
class DatasetProfile:
    """Profile for an entire dataset."""

    dataset_version: str
    source_variant: str
    workbooks: tuple[WorkbookProfile, ...]
    measurement_kind: str = "UNKNOWN"
    unit: str = "UNKNOWN"

    def to_json_string(self) -> str:
        """Serialize profile to canonical JSON string."""
        data = {
            "dataset_version": self.dataset_version,
            "source_variant": self.source_variant,
            "workbooks": [
                {
                    "relative_path": wb.relative_path,
                    "sheets": [
                        {
                            "sheet_title": s.sheet_title,
                            "max_row": s.max_row,
                            "max_column": s.max_column,
                            "header_texts": list(s.header_texts),
                            "null_count": s.null_count,
                            "datetime_count": s.datetime_count,
                            "date_count": s.date_count,
                            "time_count": s.time_count,
                            "numeric_cell_count": s.numeric_cell_count,
                            "numeric_text_count": s.numeric_text_count,
                            "iso_timestamp_text_count": s.iso_timestamp_text_count,
                            "other_text_count": s.other_text_count,
                        }
                        for s in wb.sheets
                    ],
                }
                for wb in self.workbooks
            ],
            "measurement_kind": self.measurement_kind,
            "unit": self.unit,
        }
        content = json.dumps(data, sort_keys=True, indent=2, ensure_ascii=True)
        content += "\n"
        return content

    def to_json(self, path: Path) -> None:
        """Write profile to JSON file atomically."""
        path.parent.mkdir(parents=True, exist_ok=True)
        content = self.to_json_string()

        fd, tmp_path = tempfile.mkstemp(
            dir=path.parent,
            prefix=path.name + ".",
            suffix=".tmp",
        )
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as f:
                f.write(content)
            os.replace(tmp_path, path)
        except Exception:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass
            raise


def _is_iso_timestamp_text(value: str) -> bool:
    """Check if string is an ISO-8601 timestamp."""
    return bool(_ISO_TIMESTAMP_PATTERN.match(value))


def _is_numeric_text(value: str) -> bool:
    """Check if string can be parsed as Decimal."""
    try:
        Decimal(value)
        return True
    except (InvalidOperation, ValueError):
        return False


def _compute_sha256(path: Path) -> str:
    """Compute SHA-256 hash of a file using streaming."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _classify_datetime_cell(cell: Cell) -> tuple[int, int, int]:
    """Classify a datetime cell based on its number format.

    Returns (datetime_count, date_count, time_count).
    Uses openpyxl.styles.numbers.is_date_format for classification.
    Midnight datetime is still datetime, not date.
    """
    from datetime import timedelta

    value = cell.value
    number_format = getattr(cell, "number_format", "")

    if isinstance(value, datetime):
        # Use is_date_format to check if it's a date format
        if is_date_format(number_format):
            # Check if format contains time components
            fmt_lower = number_format.lower()
            has_time = any(
                code in fmt_lower
                for code in ["h:", "h ", ":mm", ":ss", "hh:", "hh ", "[h", "[m", "[s"]
            )
            if has_time:
                return (1, 0, 0)  # datetime
            else:
                return (0, 1, 0)  # date
        else:
            # Not a recognized date format, treat as datetime
            return (1, 0, 0)

    elif isinstance(value, time):
        return (0, 0, 1)

    elif isinstance(value, date):
        return (0, 1, 0)

    elif isinstance(value, timedelta):
        # Timedelta is used for elapsed time formats - treat as datetime
        return (1, 0, 0)

    return (0, 0, 0)


def _profile_cell(cell: Cell) -> dict[str, int]:
    """Profile a single cell and return count updates."""
    from datetime import timedelta

    counts = {
        "null_count": 0,
        "datetime_count": 0,
        "date_count": 0,
        "time_count": 0,
        "numeric_cell_count": 0,
        "numeric_text_count": 0,
        "iso_timestamp_text_count": 0,
        "other_text_count": 0,
    }

    if cell.value is None:
        counts["null_count"] = 1
        return counts

    value = cell.value

    if isinstance(value, datetime):
        dt_c, d_c, t_c = _classify_datetime_cell(cell)
        counts["datetime_count"] = dt_c
        counts["date_count"] = d_c
        counts["time_count"] = t_c
        return counts
    elif isinstance(value, time):
        counts["time_count"] = 1
        return counts
    elif isinstance(value, date):
        counts["date_count"] = 1
        return counts
    elif isinstance(value, timedelta):
        # Timedelta is used for elapsed time formats - classify as datetime
        counts["datetime_count"] = 1
        return counts
    elif isinstance(value, bool):
        # Booleans not counted as numeric or any other category
        return counts
    elif isinstance(value, (int, float)):
        counts["numeric_cell_count"] = 1
        return counts
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            counts["null_count"] = 1
            return counts

        if _is_iso_timestamp_text(text):
            counts["iso_timestamp_text_count"] = 1
            return counts

        if _is_numeric_text(text):
            counts["numeric_text_count"] = 1
            return counts

        counts["other_text_count"] = 1
        return counts

    counts["other_text_count"] = 1
    return counts


def _profile_sheet(worksheet: Worksheet) -> SheetProfile:
    """Profile a single worksheet using one-pass iter_rows streaming."""
    sheet_title = worksheet.title
    max_row = worksheet.max_row
    max_column = worksheet.max_column

    # Extract header texts from row 1
    header_texts: list[str] = []

    null_count = 0
    datetime_count = 0
    date_count = 0
    time_count = 0
    numeric_cell_count = 0
    numeric_text_count = 0
    iso_timestamp_text_count = 0
    other_text_count = 0

    # Use iter_rows for one-pass streaming
    for row_idx, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_row), start=1
    ):
        for cell in row:
            if row_idx == 1:
                # Header row - capture texts
                header_value = cell.value
                if header_value is None:
                    header_texts.append("")
                else:
                    header_texts.append(str(header_value))
            else:
                # Data rows - profile cells
                counts = _profile_cell(cell)
                null_count += counts["null_count"]
                datetime_count += counts["datetime_count"]
                date_count += counts["date_count"]
                time_count += counts["time_count"]
                numeric_cell_count += counts["numeric_cell_count"]
                numeric_text_count += counts["numeric_text_count"]
                iso_timestamp_text_count += counts["iso_timestamp_text_count"]
                other_text_count += counts["other_text_count"]

    return SheetProfile(
        sheet_title=sheet_title,
        max_row=max_row,
        max_column=max_column,
        header_texts=tuple(header_texts),
        null_count=null_count,
        datetime_count=datetime_count,
        date_count=date_count,
        time_count=time_count,
        numeric_cell_count=numeric_cell_count,
        numeric_text_count=numeric_text_count,
        iso_timestamp_text_count=iso_timestamp_text_count,
        other_text_count=other_text_count,
    )


def _profile_workbook(
    file_manifest: FileManifest,
    root: Path,
) -> WorkbookProfile:
    """Profile a single workbook file."""
    file_path = root / file_manifest.relative_path

    if not file_path.exists():
        raise ValueError(f"Registered file missing: {file_manifest.relative_path}")

    actual_size = file_path.stat().st_size
    if actual_size != file_manifest.byte_size:
        raise ValueError(
            f"File size changed for {file_manifest.relative_path}: "
            f"expected {file_manifest.byte_size}, got {actual_size}"
        )

    actual_hash = _compute_sha256(file_path)
    if actual_hash != file_manifest.sha256:
        raise ValueError(f"File hash changed for {file_manifest.relative_path}")

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        raise RuntimeError(
            f"Failed to open workbook {file_manifest.relative_path}: {e}"
        ) from e

    try:
        sheet_profiles: list[SheetProfile] = []
        seen_titles: set[str] = set()

        for worksheet in wb.worksheets:
            if worksheet.title in seen_titles:
                raise ValueError(
                    f"Duplicate sheet title '{worksheet.title}' in "
                    f"{file_manifest.relative_path}"
                )
            seen_titles.add(worksheet.title)

            sheet_profiles.append(_profile_sheet(worksheet))
    finally:
        wb.close()

    return WorkbookProfile(
        relative_path=file_manifest.relative_path,
        sheets=tuple(sheet_profiles),
    )


def profile_dataset(
    manifest: DatasetManifest,
    output: Path,
) -> DatasetProfile:
    """Profile a registered dataset.

    Args:
        manifest: Dataset manifest with file registrations.
        output: Path to write profile JSON.

    Returns:
        DatasetProfile with workbook statistics.

    Raises:
        ValueError: If validation fails.
        RuntimeError: If workbook cannot be read.
    """
    root = manifest._root
    if root is None:
        raise ValueError("Manifest has no root path for profiling")

    workbook_profiles: list[WorkbookProfile] = []

    for file_manifest in manifest.files:
        if file_manifest.media_kind == "workbook":
            workbook_profiles.append(_profile_workbook(file_manifest, root))

    profile = DatasetProfile(
        dataset_version=manifest.dataset_version,
        source_variant=manifest.source_variant,
        workbooks=tuple(workbook_profiles),
        measurement_kind="UNKNOWN",
        unit="UNKNOWN",
    )

    profile.to_json(output)

    return profile
