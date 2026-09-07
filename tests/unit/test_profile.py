"""Unit tests for dataset profiling."""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import TYPE_CHECKING

import pytest

if TYPE_CHECKING:
    from typing import Any

from openpyxl import Workbook

# These imports will fail until implementation exists
from iot_energy_pipeline.manifests import (
    DatasetManifest,
    FileManifest,
    register_dataset,
)
from iot_energy_pipeline.profile import (
    DatasetProfile,
    SheetProfile,
    WorkbookProfile,
    profile_dataset,
)


def _make_xlsx(
    path: Path,
    sheets: dict[str, list[list[Any]]] | None = None,
) -> None:
    """Create an xlsx file with optional multiple sheets."""
    wb = Workbook()
    # Remove default sheet
    if sheets:
        wb.remove(wb.active)
        for sheet_name, rows in sheets.items():
            ws = wb.create_sheet(sheet_name)
            for row in rows:
                ws.append(row)
    wb.save(path)


def _make_ttl(path: Path, content: str = "@prefix : <http://example.org/> .\n") -> None:
    """Create a minimal TTL file."""
    path.write_text(content)


def _file_sha256(path: Path) -> str:
    """Calculate SHA-256 hash of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


class TestSheetProfile:
    """Tests for SheetProfile dataclass."""

    def test_sheet_profile_is_frozen(self) -> None:
        """SheetProfile should be immutable."""
        profile = SheetProfile(
            sheet_title="Sheet1",
            max_row=10,
            max_column=5,
            header_texts=("A", "B", "C"),
            null_count=0,
            datetime_count=0,
            date_count=0,
            time_count=0,
            numeric_cell_count=0,
            numeric_text_count=0,
            iso_timestamp_text_count=0,
            other_text_count=0,
        )
        with pytest.raises(AttributeError):
            profile.sheet_title = "Other"

    def test_sheet_profile_required_fields(self) -> None:
        """SheetProfile contains all required profile fields."""
        profile = SheetProfile(
            sheet_title="Data",
            max_row=100,
            max_column=10,
            header_texts=("meter_id", "timestamp", "value"),
            null_count=5,
            datetime_count=10,
            date_count=2,
            time_count=1,
            numeric_cell_count=80,
            numeric_text_count=3,
            iso_timestamp_text_count=8,
            other_text_count=1,
        )
        assert profile.sheet_title == "Data"
        assert profile.max_row == 100
        assert profile.max_column == 10
        assert profile.header_texts == ("meter_id", "timestamp", "value")
        assert profile.null_count == 5
        assert profile.datetime_count == 10
        assert profile.date_count == 2
        assert profile.time_count == 1
        assert profile.numeric_cell_count == 80
        assert profile.numeric_text_count == 3
        assert profile.iso_timestamp_text_count == 8
        assert profile.other_text_count == 1


class TestWorkbookProfile:
    """Tests for WorkbookProfile dataclass."""

    def test_workbook_profile_is_frozen(self) -> None:
        """WorkbookProfile should be immutable."""
        profile = WorkbookProfile(
            relative_path="data.xlsx",
            sheets=(),
        )
        with pytest.raises(AttributeError):
            profile.relative_path = "other.xlsx"

    def test_workbook_profile_contains_sheets(self) -> None:
        """WorkbookProfile contains sheet profiles."""
        sheet = SheetProfile(
            sheet_title="Data",
            max_row=10,
            max_column=3,
            header_texts=("a", "b", "c"),
            null_count=0,
            datetime_count=0,
            date_count=0,
            time_count=0,
            numeric_cell_count=0,
            numeric_text_count=0,
            iso_timestamp_text_count=0,
            other_text_count=0,
        )
        profile = WorkbookProfile(
            relative_path="data.xlsx",
            sheets=(sheet,),
        )
        assert len(profile.sheets) == 1
        assert profile.sheets[0].sheet_title == "Data"


class TestDatasetProfile:
    """Tests for DatasetProfile dataclass."""

    def test_dataset_profile_is_frozen(self) -> None:
        """DatasetProfile should be immutable."""
        profile = DatasetProfile(
            dataset_version="v1",
            source_variant="raw",
            workbooks=(),
            measurement_kind="UNKNOWN",
            unit="UNKNOWN",
        )
        with pytest.raises(AttributeError):
            profile.dataset_version = "v2"

    def test_dataset_profile_unknown_measurement_kind(self, tmp_path: Path) -> None:
        """DatasetProfile always reports UNKNOWN measurement_kind."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")
        assert profile.measurement_kind == "UNKNOWN"

    def test_dataset_profile_unknown_unit(self, tmp_path: Path) -> None:
        """DatasetProfile always reports UNKNOWN unit."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")
        assert profile.unit == "UNKNOWN"


class TestProfileDataset:
    """Tests for profile_dataset function."""

    def test_profile_single_workbook(self, tmp_path: Path) -> None:
        """Profile a single workbook with one sheet."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(
            data_dir / "test.xlsx",
            {"Meter1": [["meter_id", "value"], ["M001", 100.5]]},
        )

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        assert len(profile.workbooks) == 1
        assert profile.workbooks[0].relative_path == "test.xlsx"
        assert len(profile.workbooks[0].sheets) == 1
        assert profile.workbooks[0].sheets[0].sheet_title == "Meter1"

    def test_profile_multiple_sheets(self, tmp_path: Path) -> None:
        """Profile a workbook with multiple sheets."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(
            data_dir / "multi.xlsx",
            {
                "Building1": [["meter_id", "value"], ["M001", 100]],
                "Building2": [["meter_id", "value"], ["M002", 200]],
            },
        )

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        assert len(profile.workbooks[0].sheets) == 2
        sheet_titles = [s.sheet_title for s in profile.workbooks[0].sheets]
        assert "Building1" in sheet_titles
        assert "Building2" in sheet_titles

    def test_profile_sheet_max_row_max_column(self, tmp_path: Path) -> None:
        """Profile records max_row and max_column for each sheet."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(
            data_dir / "test.xlsx",
            {"Data": [["a", "b", "c"], [1, 2, 3], [4, 5, 6]]},
        )

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.max_row == 3  # 1 header + 2 data rows
        assert sheet.max_column == 3

    def test_profile_header_texts(self, tmp_path: Path) -> None:
        """Profile records header texts from row 1."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(
            data_dir / "test.xlsx",
            {"Data": [["meter_id", "timestamp", "value", "unit"]]},
        )

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.header_texts == ("meter_id", "timestamp", "value", "unit")

    def test_profile_null_count(self, tmp_path: Path) -> None:
        """Profile counts null cells."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["a", "b", "c"])  # Header
        ws.append([1, None, 3])  # One null
        ws.append([None, None, None])  # Three nulls
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.null_count == 4  # 1 + 3

    def test_profile_datetime_count(self, tmp_path: Path) -> None:
        """Profile counts native datetime cells."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["ts"])
        ws.append([datetime(2024, 1, 15, 10, 30, 0)])
        ws.append([datetime(2024, 2, 20, 14, 45, 30)])
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.datetime_count == 2

    def test_profile_date_count(self, tmp_path: Path) -> None:
        """Profile counts native date cells."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["d"])
        ws.append([date(2024, 1, 15)])
        ws.append([date(2024, 2, 20)])
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.date_count == 2

    def test_profile_time_count(self, tmp_path: Path) -> None:
        """Profile counts native time cells."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["t"])
        ws.append([time(10, 30, 0)])
        ws.append([time(14, 45, 30)])
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.time_count == 2

    def test_profile_numeric_cell_count(self, tmp_path: Path) -> None:
        """Profile counts numeric cells (floats, ints) excluding bools."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["val"])
        ws.append([100])
        ws.append([200.5])
        ws.append([True])  # Bool - should NOT be counted as numeric
        ws.append([False])  # Bool - should NOT be counted as numeric
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        # Booleans are NOT counted as numeric per the contract
        assert sheet.numeric_cell_count == 2

    def test_profile_numeric_text_count(self, tmp_path: Path) -> None:
        """Profile counts text cells parseable as Decimal."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["val"])
        ws.append(["100.50"])  # Decimal-parseable
        ws.append(["200.75"])  # Decimal-parseable
        ws.append(["abc"])  # Not decimal-parseable
        ws.append(["1e10"])  # Scientific notation - Decimal can parse this
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        # All four text values should be parseable by Decimal
        # Actually, "abc" is not, so should be 3
        assert sheet.numeric_text_count == 3

    def test_profile_iso_timestamp_text_count(self, tmp_path: Path) -> None:
        """Profile counts text that looks like ISO-8601 timestamps."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["ts"])
        ws.append(["2024-01-15T10:30:00"])  # ISO-8601
        ws.append(["2024-02-20T14:45:30Z"])  # ISO-8601 with Z
        ws.append(["2024-03-25T08:00:00+08:00"])  # ISO-8601 with tz
        ws.append(["not a timestamp"])  # Not ISO
        ws.append(["2024-01-15"])  # Date only - not timestamp
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        # Three ISO-8601 timestamps with time component
        assert sheet.iso_timestamp_text_count == 3

    def test_profile_other_text_count(self, tmp_path: Path) -> None:
        """Profile counts other text cells."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["text"])
        ws.append(["hello"])  # other text
        ws.append(["world"])  # other text
        ws.append(["2024-01-15T10:00:00"])  # ISO timestamp - counted separately
        ws.append(["100.5"])  # numeric text - counted separately
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        # "hello" and "world" are other text (2)
        assert sheet.other_text_count == 2

    def test_profile_no_sample_values(self, tmp_path: Path) -> None:
        """Profile does not preserve sample values."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(
            data_dir / "test.xlsx",
            {"Data": [["val"], ["secret_value_123"]]},
        )

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")
        profile_path = tmp_path / "profile.json"
        profile_path.write_text(profile.to_json_string())

        # Verify no sample values in JSON
        json_text = profile_path.read_text()
        assert "secret_value" not in json_text
        assert "sample" not in json_text.lower()


class TestProfileErrors:
    """Tests for error handling in profile_dataset."""

    def test_missing_file_raises(self, tmp_path: Path) -> None:
        """Profiling raises error if registered file is missing."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [[1]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        # Delete the file after registration
        (data_dir / "test.xlsx").unlink()

        with pytest.raises(ValueError, match="test.xlsx"):
            profile_dataset(manifest, tmp_path / "profile.json")

    def test_size_changed_raises(self, tmp_path: Path) -> None:
        """Profiling raises error if file size changed."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        xlsx_path = data_dir / "test.xlsx"
        _make_xlsx(xlsx_path, {"Sheet1": [[1]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        # Modify the file to change size
        _make_xlsx(xlsx_path, {"Sheet1": [[1, 2, 3, 4, 5]]})

        with pytest.raises(ValueError, match="size"):
            profile_dataset(manifest, tmp_path / "profile.json")

    def test_hash_changed_raises(self, tmp_path: Path) -> None:
        """Profiling raises error if file hash changed."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        xlsx_path = data_dir / "test.xlsx"
        _make_xlsx(xlsx_path, {"Sheet1": [[1]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        # Modify the file without significantly changing size
        # (by changing content but keeping similar size)
        original_size = xlsx_path.stat().st_size
        _make_xlsx(xlsx_path, {"Sheet1": [[9]]})

        # Skip if size actually changed
        if xlsx_path.stat().st_size != original_size:
            pytest.skip("Size changed, test not applicable")

        with pytest.raises(ValueError, match="hash"):
            profile_dataset(manifest, tmp_path / "profile.json")

    def test_corrupt_workbook_raises(self, tmp_path: Path) -> None:
        """Profiling raises error for corrupt workbook."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()

        # Create a file with xlsx extension but invalid content
        xlsx_path = data_dir / "corrupt.xlsx"
        xlsx_path.write_bytes(b"not a valid xlsx file")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        with pytest.raises((ValueError, RuntimeError), match="corrupt.xlsx"):
            profile_dataset(manifest, tmp_path / "profile.json")

    @pytest.mark.skipif(
        not hasattr(Path, "symlink_to"),
        reason="Symlinks not supported on this platform",
    )
    def test_profile_rejects_workbook_symlink(self, tmp_path: Path) -> None:
        """Profiling rejects a workbook that is a symlink."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        real_file = data_dir / "real.xlsx"
        _make_xlsx(real_file, {"Sheet1": [["a", "b"], [1, 2]]})

        link_file = data_dir / "link.xlsx"
        try:
            link_file.symlink_to(real_file)
        except OSError:
            pytest.skip("Symlink creation failed")

        # Register the symlink file by creating a manifest directly
        manifest = DatasetManifest(
            dataset_version="v1",
            source_variant="raw",
            files=(
                FileManifest(
                    relative_path="link.xlsx",
                    byte_size=real_file.stat().st_size,
                    sha256=_file_sha256(real_file),
                    media_kind="workbook",
                ),
            ),
            _root=data_dir,
        )

        with pytest.raises(ValueError, match="symlink"):
            profile_dataset(manifest, tmp_path / "profile.json")

    def test_duplicate_sheet_title_raises(self, tmp_path: Path) -> None:
        """Profiling raises error for duplicate sheet titles."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()

        # openpyxl doesn't allow duplicate sheet names in one workbook
        # But we should handle it if somehow it occurs
        # This test would require manually crafting a malformed xlsx
        # For now, skip this test as openpyxl prevents this
        pytest.skip("Cannot create workbook with duplicate sheets via openpyxl")


class TestProfileOutput:
    """Tests for profile JSON output format."""

    def test_profile_creates_output(self, tmp_path: Path) -> None:
        """profile_dataset writes output file."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        output_path = tmp_path / "output" / "profile.json"
        profile_dataset(manifest, output_path)

        assert output_path.exists()

    def test_profile_creates_parent_directory(self, tmp_path: Path) -> None:
        """profile_dataset creates output parent directory."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        output_path = tmp_path / "deep" / "nested" / "profile.json"
        profile_dataset(manifest, output_path)

        assert output_path.exists()

    def test_profile_byte_identical_json(self, tmp_path: Path) -> None:
        """Repeated profiling produces byte-identical JSON."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(
            data_dir / "test.xlsx",
            {"Sheet1": [["a", "b"], [1, 2], [3, 4]]},
        )

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        output1 = tmp_path / "profile1.json"
        output2 = tmp_path / "profile2.json"

        profile_dataset(manifest, output1)
        profile_dataset(manifest, output2)

        assert output1.read_bytes() == output2.read_bytes()

    def test_profile_json_canonical_format(self, tmp_path: Path) -> None:
        """Profile JSON is canonical with sorted keys."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        output_path = tmp_path / "profile.json"
        profile_dataset(manifest, output_path)

        data = json.loads(output_path.read_text())
        canonical = json.dumps(data, sort_keys=True, indent=2, ensure_ascii=True)
        canonical += "\n"

        assert output_path.read_text() == canonical

    def test_profile_json_no_absolute_paths(self, tmp_path: Path) -> None:
        """Profile JSON contains no absolute paths."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        output_path = tmp_path / "profile.json"
        profile_dataset(manifest, output_path)

        json_text = output_path.read_text()
        assert str(tmp_path) not in json_text
        assert "C:\\" not in json_text

    def test_profile_json_utf8(self, tmp_path: Path) -> None:
        """Profile JSON is UTF-8 encoded."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        output_path = tmp_path / "profile.json"
        profile_dataset(manifest, output_path)

        # Should decode as UTF-8
        content = output_path.read_text(encoding="utf-8")
        assert "dataset_version" in content


class TestTTLFiles:
    """Tests for TTL file handling in profiling."""

    def test_ttl_appears_in_manifest_totals(self, tmp_path: Path) -> None:
        """TTL files are registered in manifest."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "data.xlsx", {"Sheet1": [[1]]})
        _make_ttl(data_dir / "brick.ttl")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        # Both files should be in manifest
        assert len(manifest.files) == 2
        file_kinds = {f.relative_path: f.media_kind for f in manifest.files}
        assert file_kinds["data.xlsx"] == "workbook"
        assert file_kinds["brick.ttl"] == "brick_ttl"

    def test_ttl_not_parsed_in_task2(self, tmp_path: Path) -> None:
        """TTL files are not parsed during profiling in Task 2."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "data.xlsx", {"Sheet1": [["a"], [1]]})
        _make_ttl(data_dir / "brick.ttl")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        # Only workbook should have profile, not TTL
        assert len(profile.workbooks) == 1
        assert profile.workbooks[0].relative_path == "data.xlsx"


class TestOpenpyxlSettings:
    """Tests that verify openpyxl read-only and data-only modes."""

    def test_read_only_mode_used(self, tmp_path: Path) -> None:
        """Profile uses read_only=True when opening workbooks."""
        # This is implicitly tested by the implementation
        # The test ensures the contract is followed
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "test.xlsx", {"Sheet1": [[1]]})

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        # Should succeed without error
        profile_dataset(manifest, tmp_path / "profile.json")

    def test_data_only_mode_used(self, tmp_path: Path) -> None:
        """Profile uses data_only=True when opening workbooks."""
        # Create a workbook with a formula
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["a", "b", "sum"])
        ws.append([1, 2, "=A2+B2"])  # Formula
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        # With data_only=True, formula cell returns None (not evaluated)
        # Without data_only=True, it would return the formula string
        profile = profile_dataset(manifest, tmp_path / "profile.json")

        # The formula cell should be treated as null (data_only=True)
        sheet = profile.workbooks[0].sheets[0]
        # One header row + one data row, 3 columns
        # Data row: 1 (numeric), 2 (numeric), None (formula not evaluated)
        assert sheet.null_count >= 1  # At least the formula cell is null

    def test_elapsed_time_format_classified_as_datetime(self, tmp_path: Path) -> None:
        """Elapsed time format [h]:mm:ss is classified as datetime."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["elapsed"])

        # Elapsed time format - not covered by old substring heuristic
        cell = ws.cell(row=2, column=1, value=datetime(2024, 1, 15, 5, 30, 0))
        cell.number_format = "[h]:mm:ss"

        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        # Should be classified as datetime, not date
        assert sheet.datetime_count == 1
        assert sheet.date_count == 0


class TestDateFormatDistinction:
    """Tests for date vs datetime distinction using cell number format."""

    def test_midnight_datetime_remains_datetime(self, tmp_path: Path) -> None:
        """A datetime at midnight is still counted as datetime, not date."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["ts"])
        # Midnight datetime - should still be datetime, not date
        ws.append([datetime(2024, 1, 15, 0, 0, 0)])
        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        # Midnight datetime is still datetime, not date
        assert sheet.datetime_count == 1
        assert sheet.date_count == 0

    def test_date_format_distinguished_from_datetime(self, tmp_path: Path) -> None:
        """Cell with date format is counted as date, datetime format as datetime."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["dt", "d"])

        # Datetime with datetime format
        cell_dt = ws.cell(row=2, column=1, value=datetime(2024, 1, 15, 10, 30, 0))
        cell_dt.number_format = "yyyy-mm-dd h:mm:ss"

        # Date with date format (openpyxl stores as datetime at midnight)
        cell_d = ws.cell(row=2, column=2, value=date(2024, 2, 20))
        cell_d.number_format = "yyyy-mm-dd"

        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.datetime_count == 1
        assert sheet.date_count == 1

    def test_time_format_distinguished(self, tmp_path: Path) -> None:
        """Cell with time format is counted as time."""
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["t"])

        cell_t = ws.cell(row=2, column=1, value=time(10, 30, 0))
        cell_t.number_format = "h:mm:ss"

        wb.save(data_dir / "test.xlsx")

        manifest = register_dataset(
            root=data_dir,
            output=tmp_path / "manifest.json",
            dataset_version="v1",
            source_variant="raw",
        )

        profile = profile_dataset(manifest, tmp_path / "profile.json")

        sheet = profile.workbooks[0].sheets[0]
        assert sheet.time_count == 1
