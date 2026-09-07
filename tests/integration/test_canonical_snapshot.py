from __future__ import annotations

import hashlib
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import iot_energy_pipeline.canonicalize as canonicalize_module
from iot_energy_pipeline.canonicalize import (
    CanonicalizationRules,
    MeasurementRule,
    build_snapshot,
    canonicalize_workbook,
)
from iot_energy_pipeline.manifests import register_dataset
from iot_energy_pipeline.source_layout import (
    SourceLayoutRules,
    WorkbookLayoutRule,
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_snapshot_does_not_materialize_results_above_the_return_boundary(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    manifest = _source(tmp_path)
    monkeypatch.setattr(
        canonicalize_module, "_SNAPSHOT_RESULT_MATERIALIZATION_LIMIT", 2
    )

    result = build_snapshot(manifest, tmp_path / "snapshot", _rules())

    assert result.event_count == 3
    assert result.events is None
    assert (
        len((tmp_path / "snapshot/canonical.ndjson").read_text("utf-8").splitlines())
        == 3
    )


def _source(tmp_path: Path, *, variant: str = "clean"):
    root = tmp_path / variant
    root.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Readings"
    sheet.append(["timestamp", "meter_id", "value"])
    sheet.append([datetime(2024, 1, 1, 0, 0), "meter-a", "10.00"])
    sheet.append([datetime(2024, 1, 1, 0, 0), "meter-a", "10.00"])
    sheet.append([datetime(2024, 1, 1, 0, 15), "meter-a", None])
    sheet.append([datetime(2024, 1, 1, 0, 30), "meter-a", "not-a-number"])
    sheet.append([datetime(2024, 1, 1, 0, 45), "meter-b", "2.50"])
    workbook.save(root / "readings.xlsx")
    (root / "metadata.ttl").write_text(
        """\
@prefix brick: <https://brickschema.org/schema/Brick#> .
@prefix ref: <https://brickschema.org/schema/Brick/ref#> .
@prefix unit: <http://qudt.org/vocab/unit/> .
@prefix ex: <https://example.test/> .
ex:meter-entity a brick:Electrical_Meter ;
  brick:hasUnit unit:KiloW-HR ;
  ref:hasExternalReference [ ref:hasTimeseriesId "meter-a" ] .
ex:building-a a brick:Building ; brick:hasPart ex:zone-a .
ex:zone-a a brick:Zone ; brick:hasPart ex:meter-entity .
ex:panel-a a brick:Equipment ; brick:isMeteredBy ex:meter-entity .
""",
        encoding="utf-8",
    )
    return register_dataset(
        root,
        tmp_path / f"{variant}-dataset-manifest.json",
        dataset_version="fixture-v1",
        source_variant=variant,  # type: ignore[arg-type]
    )


def _rules() -> CanonicalizationRules:
    evidence = "synthetic fixture declaration"
    return CanonicalizationRules(
        source_timezone="Asia/Hong_Kong",
        timezone_evidence=evidence,
        source_layout=SourceLayoutRules(
            rules=(
                WorkbookLayoutRule(
                    name="synthetic-long-format",
                    relative_path_pattern=r"readings\.xlsx",
                    sheet_name="Readings",
                    header_row=1,
                    expected_columns=("timestamp", "meter_id", "value"),
                    timestamp_column="timestamp",
                    value_column="value",
                    meter_id_column="meter_id",
                    evidence=evidence,
                ),
            ),
            evidence=evidence,
        ),
        measurements={
            "meter-a": MeasurementRule(
                "cumulative_energy", "kWh", 2, evidence, "KiloW-HR"
            ),
            "meter-b": MeasurementRule("cumulative_energy", "kWh", 2, evidence, None),
        },
    )


def test_canonicalization_requires_declared_timezone_and_measurement_evidence(
    tmp_path: Path,
) -> None:
    manifest = _source(tmp_path)

    with pytest.raises(ValueError, match="canonicalization rules"):
        canonicalize_workbook(manifest)


def test_snapshot_preserves_rows_and_reports_invalid_source_rows(
    tmp_path: Path,
) -> None:
    manifest = _source(tmp_path)

    result = build_snapshot(manifest, tmp_path / "snapshot", _rules())

    assert len(result.events) == 3
    assert len({event["source_event_id"] for event in result.events}) == 3
    duplicate_time = [
        event
        for event in result.events
        if event["meter_id"] == "meter-a"
        and event["original_timestamp_text"].startswith("2024-01-01T00:00:00")
    ]
    assert len(duplicate_time) == 2
    assert {row["category"] for row in result.source_quality} == {
        "missing_numeric_value",
        "invalid_numeric_value",
    }
    unmatched = next(e for e in result.events if e["meter_id"] == "meter-b")
    assert unmatched["brick_entity_id"] is None
    assert unmatched["quality_flags"] == ["brick_unmatched"]
    matched = next(e for e in result.events if e["meter_id"] == "meter-a")
    assert matched["brick_building_ids"] == ["https://example.test/building-a"]
    assert matched["brick_zone_ids"] == ["https://example.test/zone-a"]
    assert matched["brick_metered_entities"] == [
        {
            "entity_id": "https://example.test/panel-a",
            "type_uris": ["https://brickschema.org/schema/Brick#Equipment"],
        }
    ]
    assert matched["brick_unit_uri"].endswith("KiloW-HR")
    assert matched["event_time_utc"].endswith("Z")
    assert matched["event_time_utc"].startswith("2023-12-31T16:00:00")
    assert matched["original_timestamp_text"].startswith("2024-01-01T00:00:00")
    assert matched["original_numeric_text"] == "10.00"


def test_snapshot_preserves_subsecond_source_precision_and_uses_sortable_utc(
    tmp_path: Path,
) -> None:
    manifest = _source(tmp_path)
    workbook_path = manifest._root / "readings.xlsx"
    workbook = load_workbook(workbook_path)
    workbook["Readings"]["A2"] = datetime(2024, 1, 1, 0, 0, 0, 500000)
    workbook.save(workbook_path)
    updated = register_dataset(
        manifest._root,
        tmp_path / "fractional-dataset-manifest.json",
        dataset_version="fixture-v1",
        source_variant="clean",
    )

    result = build_snapshot(updated, tmp_path / "snapshot", _rules())

    row = next(event for event in result.events if event["source_row_number"] == 2)
    assert row["original_timestamp_text"] == "2024-01-01T00:00:00.500000"
    assert row["event_time_utc"] == "2023-12-31T16:00:00.500000Z"


def test_ndjson_and_parquet_contain_the_same_sorted_exact_records(
    tmp_path: Path,
) -> None:
    manifest = _source(tmp_path)
    output = tmp_path / "snapshot"

    result = build_snapshot(manifest, output, _rules())

    ndjson_records = [
        json.loads(line)
        for line in (
            (output / "canonical.ndjson").read_text(encoding="utf-8").splitlines()
        )
    ]
    import pyarrow.parquet as pq

    parquet_records = pq.read_table(output / "canonical.parquet").to_pylist()
    assert ndjson_records == parquet_records == list(result.events)
    assert [row["source_event_id"] for row in ndjson_records] == sorted(
        row["source_event_id"] for row in ndjson_records
    )
    assert (
        pq.read_schema(output / "canonical.parquet").field("scaled_value").type
        == __import__("pyarrow").int64()
    )


def test_snapshot_manifest_binds_every_input_and_output_hash(tmp_path: Path) -> None:
    manifest = _source(tmp_path)
    output = tmp_path / "snapshot"

    result = build_snapshot(manifest, output, _rules())
    saved = json.loads((output / "snapshot-manifest.json").read_text("utf-8"))

    assert saved == result.manifest
    assert saved["status"] == "complete"
    assert saved["source_variant"] == "clean"
    assert saved["source_locator_fields"] == [
        "dataset_version",
        "source_variant",
        "source_relative_path",
        "workbook_sheet",
        "source_row_number",
    ]
    assert saved["artifacts"]["canonical_ndjson"]["sha256"] == _sha256(
        output / "canonical.ndjson"
    )
    assert saved["artifacts"]["canonical_parquet"]["sha256"] == _sha256(
        output / "canonical.parquet"
    )
    assert saved["brick"]["sha256"] == next(
        file.sha256 for file in manifest.files if file.media_kind == "brick_ttl"
    )
    assert saved["measurement_registry"]["meter-a"]["evidence"]
    assert saved["timezone"]["evidence"]
    assert saved["source_layout"]["sha256"] == _rules().source_layout.sha256()


def test_per_workbook_hkust_layout_derives_meter_from_registered_path(
    tmp_path: Path,
) -> None:
    root = tmp_path / "raw"
    source_dir = root / "Time-series data"
    source_dir.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["time", "number"])
    sheet.append([datetime(2024, 1, 1, 0, 0), "10.10"])
    workbook.save(source_dir / "GUI.NO.D0001.xlsx")
    (root / "HKUST_Meter_Metadata.ttl").write_text(
        """\
@prefix brick: <https://brickschema.org/schema/Brick#> .
@prefix ref: <https://brickschema.org/schema/Brick/ref#> .
@prefix unit: <http://qudt.org/vocab/unit/> .
@prefix ex: <https://example.test/> .
ex:meter a brick:Electrical_Meter ;
  brick:hasUnit unit:KiloW-HR ;
  ref:hasExternalReference [ ref:hasTimeseriesId "D0001" ] .
""",
        encoding="utf-8",
    )
    manifest = register_dataset(
        root,
        tmp_path / "manifest.json",
        dataset_version="dryad-v5",
        source_variant="raw",
    )
    evidence = "synthetic representation of the authors' documented source layout"
    rules = CanonicalizationRules(
        source_timezone="Asia/Hong_Kong",
        timezone_evidence=evidence,
        source_layout=SourceLayoutRules(
            rules=(
                WorkbookLayoutRule(
                    name="hkust-raw-per-meter",
                    relative_path_pattern=(
                        r"Time-series data/GUI\.NO\.(?P<meter_id>[A-Z][0-9]{4})"
                        r"\.xlsx"
                    ),
                    sheet_name="Sheet1",
                    header_row=1,
                    expected_columns=("time", "number"),
                    timestamp_column="time",
                    value_column="number",
                    meter_id_path_group="meter_id",
                    evidence=evidence,
                ),
            ),
            evidence=evidence,
        ),
        measurements={
            "D0001": MeasurementRule(
                "cumulative_energy", "kWh", 2, evidence, "KiloW-HR"
            )
        },
    )

    result = build_snapshot(manifest, tmp_path / "snapshot", rules)

    assert len(result.events) == 1
    assert result.events[0]["meter_id"] == "D0001"
    assert result.events[0]["source_relative_path"] == (
        "Time-series data/GUI.NO.D0001.xlsx"
    )


def test_unexpected_sheet_or_header_is_rejected_by_layout(tmp_path: Path) -> None:
    manifest = _source(tmp_path)
    rules = _rules()
    source = manifest._root / "readings.xlsx"  # type: ignore[operator]
    workbook = __import__("openpyxl").load_workbook(source)
    workbook.create_sheet("Undeclared")
    workbook.save(source)
    changed = register_dataset(
        manifest._root,  # type: ignore[arg-type]
        tmp_path / "changed-manifest.json",
        dataset_version="fixture-v1",
        source_variant="clean",
    )

    with pytest.raises(ValueError, match="unexpected workbook sheets"):
        build_snapshot(changed, tmp_path / "snapshot", rules)


def test_snapshot_is_immutable_and_reproducible_across_directories(
    tmp_path: Path,
) -> None:
    manifest = _source(tmp_path)
    first = tmp_path / "snapshot-a"
    second = tmp_path / "snapshot-b"

    build_snapshot(manifest, first, _rules())
    build_snapshot(manifest, second, _rules())

    for name in (
        "canonical.ndjson",
        "canonical.parquet",
        "source-quality.ndjson",
        "snapshot-manifest.json",
    ):
        assert (first / name).read_bytes() == (second / name).read_bytes()
    with pytest.raises(FileExistsError, match="immutable snapshot"):
        build_snapshot(manifest, first, _rules())


def test_contract_mismatch_between_brick_unit_and_rule_fails(tmp_path: Path) -> None:
    manifest = _source(tmp_path)
    bad_rule = _rules().with_measurement(
        "meter-a",
        MeasurementRule(
            "cumulative_energy", "kWh", 2, "synthetic fixture declaration", "W"
        ),
    )

    with pytest.raises(ValueError, match="Brick unit contract mismatch"):
        build_snapshot(manifest, tmp_path / "snapshot", bad_rule)


def test_required_brick_unit_rejects_an_unmatched_meter(tmp_path: Path) -> None:
    manifest = _source(tmp_path)
    rules = _rules().with_measurement(
        "meter-b",
        MeasurementRule(
            "cumulative_energy",
            "kWh",
            2,
            "synthetic fixture declaration",
            "KiloW-HR",
        ),
    )

    with pytest.raises(ValueError, match="Brick unit contract mismatch"):
        build_snapshot(manifest, tmp_path / "snapshot", rules)
