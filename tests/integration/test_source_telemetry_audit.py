from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from iot_energy_pipeline.manifests import register_dataset
from iot_energy_pipeline.source_audit import audit_source_telemetry
from iot_energy_pipeline.source_layout import (
    SourceLayoutRules,
    WorkbookLayoutRule,
)


def _source(tmp_path: Path):
    root = tmp_path / "raw"
    source = root / "Time-series data"
    source.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["time", "number"])
    sheet.append([datetime(2024, 1, 1, 0, 0), "10.00"])
    sheet.append([datetime(2024, 1, 1, 0, 15), "11.00"])
    sheet.append([datetime(2024, 1, 1, 0, 15), "11.50"])
    sheet.append([datetime(2024, 1, 1, 0, 45), "NaN"])
    sheet.append([datetime(2024, 1, 1, 0, 30), "9.00"])
    sheet.append([None, None])
    sheet.append(["bad-time", "bad-number"])
    sheet.append([datetime(2024, 1, 1, 1, 0), "0.00"])
    sheet.append([datetime(2024, 1, 1, 1, 15), "-1.00"])
    workbook.save(source / "GUI.NO.D0001.xlsx")
    manifest = register_dataset(
        root,
        tmp_path / "manifest.json",
        dataset_version="dryad-v5-test",
        source_variant="raw",
    )
    evidence = "synthetic source-layout test based on the authors' public code"
    layout = SourceLayoutRules(
        rules=(
            WorkbookLayoutRule(
                name="hkust-raw-per-meter",
                relative_path_pattern=(
                    r"Time-series data/GUI\.NO\.(?P<meter_id>[A-Z][0-9]{4})"
                    r"\.xlsx"
                ),
                sheet_name="Sheet1",
                header_row=1,
                expected_columns=("time", "number"),
                timestamp_column="time",
                value_column="number",
                meter_id_path_group="meter_id",
                evidence=evidence,
            ),
        ),
        evidence=evidence,
    )
    return manifest, layout


def test_audit_reports_exact_source_quality_without_assigning_semantics(
    tmp_path: Path,
) -> None:
    manifest, layout = _source(tmp_path)

    report = audit_source_telemetry(manifest, layout, tmp_path / "audit.json")

    assert report["schema_version"] == "source-telemetry-audit-v1"
    assert report["source_variant"] == "raw"
    assert report["timezone_assignment"] == "UNASSIGNED"
    assert report["measurement_semantics_assignment"] == "UNASSIGNED"
    assert report["totals"] == {
        "workbook_count": 1,
        "empty_workbook_count": 0,
        "meter_count": 1,
        "source_row_count": 9,
        "empty_source_row_count": 1,
        "missing_meter_id_count": 0,
        "missing_timestamp_count": 1,
        "invalid_timestamp_count": 1,
        "naive_timestamp_count": 7,
        "aware_timestamp_count": 0,
        "duplicate_timestamp_count": 1,
        "out_of_order_timestamp_count": 1,
        "missing_numeric_count": 1,
        "invalid_numeric_count": 1,
        "non_finite_numeric_count": 1,
        "finite_numeric_count": 6,
        "zero_numeric_count": 1,
        "negative_numeric_count": 1,
        "ordered_numeric_decrease_count": 3,
    }
    workbook = report["workbooks"][0]
    assert workbook["relative_path"] == "Time-series data/GUI.NO.D0001.xlsx"
    assert workbook["empty_workbook"] is False
    assert workbook["meter_ids"] == ["D0001"]
    assert workbook["cadence_seconds_histogram"] == {"900": 5}
    assert workbook["decimal_scale_histogram"] == {"2": 6}
    assert workbook["timestamp_bounds"] == {
        "naive": {
            "minimum": "2024-01-01T00:00:00",
            "maximum": "2024-01-01T01:15:00",
        },
        "aware_utc": {"minimum": None, "maximum": None},
    }


def test_audit_json_is_byte_reproducible_and_binds_layout(tmp_path: Path) -> None:
    manifest, layout = _source(tmp_path)
    first = tmp_path / "audit-a.json"
    second = tmp_path / "audit-b.json"

    report = audit_source_telemetry(manifest, layout, first)
    audit_source_telemetry(manifest, layout, second)

    assert first.read_bytes() == second.read_bytes()
    assert json.loads(first.read_text(encoding="utf-8")) == report
    assert report["source_layout_sha256"] == layout.sha256()
    assert report["dataset_manifest_sha256"]


def test_source_audit_is_available_through_the_cli(tmp_path: Path) -> None:
    manifest, layout = _source(tmp_path)
    manifest_path = tmp_path / "manifest.json"
    layout_path = tmp_path / "layout.json"
    output = tmp_path / "cli-audit.json"
    layout.to_json(layout_path)
    from iot_energy_pipeline.cli import main

    result = main(
        [
            "audit-source",
            "--manifest",
            str(manifest_path),
            "--root",
            str(manifest._root),
            "--layout",
            str(layout_path),
            "--output",
            str(output),
        ]
    )

    assert result == 0
    assert output.exists()


def test_valid_header_only_workbook_is_reported_as_empty(tmp_path: Path) -> None:
    manifest, layout = _source(tmp_path)
    empty = Workbook()
    sheet = empty.active
    sheet.title = "Sheet1"
    sheet.append(["time", "number"])
    empty.save(manifest._root / "Time-series data" / "GUI.NO.D0002.xlsx")
    updated = register_dataset(
        manifest._root,
        tmp_path / "updated-manifest.json",
        dataset_version="dryad-v5-test",
        source_variant="raw",
    )

    report = audit_source_telemetry(updated, layout, tmp_path / "audit.json")

    assert report["totals"]["workbook_count"] == 2
    assert report["totals"]["empty_workbook_count"] == 1
    empty_report = next(
        workbook
        for workbook in report["workbooks"]
        if workbook["relative_path"].endswith("D0002.xlsx")
    )
    assert empty_report["empty_workbook"] is True
    assert empty_report["source_row_count"] == 0
    assert empty_report["meter_ids"] == ["D0002"]


def test_source_audit_preserves_subsecond_timestamp_bounds(tmp_path: Path) -> None:
    manifest, layout = _source(tmp_path)
    workbook_path = manifest._root / "Time-series data" / "GUI.NO.D0001.xlsx"
    workbook = load_workbook(workbook_path)
    workbook["Sheet1"]["A2"] = datetime(2024, 1, 1, 0, 0, 0, 500000)
    workbook.save(workbook_path)
    updated = register_dataset(
        manifest._root,
        tmp_path / "fractional-manifest.json",
        dataset_version="dryad-v5-test",
        source_variant="raw",
    )

    report = audit_source_telemetry(updated, layout, tmp_path / "audit.json")

    assert report["workbooks"][0]["timestamp_bounds"]["naive"]["minimum"] == (
        "2024-01-01T00:00:00.500000"
    )
